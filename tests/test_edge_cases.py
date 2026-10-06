"""File-handling edge cases, extraction honesty, verifier behaviour, privacy, programs, API, backup."""
import io
import json
import re
import zipfile

import pymupdf
import pytest

from navcoach import clients, db, ingest, programs, rag
from navcoach.config import get_settings, update_settings
from navcoach.extract import ExtractionError, extract
from navcoach.llm import PrivacyBlocked, get_llm
from navcoach.verify import check_claim, locate_quote
from tests import fixtures as F


def _add(path):
    return ingest.add_file(path.name, path.read_bytes(), sync=True)


# ------------------------------------------------------------------ files

def test_empty_file_fails_honestly(env):
    p = env["files"] / "empty.txt"
    p.write_bytes(b"")
    d = _add(p)
    assert d["status"] == "failed" and "empty" in d["status_detail"]
    assert d["chunk_count"] == 0


def test_whitespace_only_file_fails(env):
    p = env["files"] / "blank.md"
    p.write_text("   \n\n  \n", encoding="utf-8")
    d = _add(p)
    assert d["status"] == "failed" and "no extractable text" in d["status_detail"]


def test_duplicate_file_is_detected(env):
    p = F.study_a(env["files"])
    _add(p)
    with pytest.raises(ingest.DuplicateFile):
        ingest.add_file("renamed.pdf", p.read_bytes(), sync=True)
    assert len(ingest.list_documents()) == 1


def test_corrupt_pdf_fails_without_crash(env):
    p = env["files"] / "broken.pdf"
    p.write_bytes(b"%PDF-1.4\n" + b"\x00garbage" * 200)
    d = _add(p)
    assert d["status"] == "failed"
    assert "could not" in d["status_detail"].lower() or "no extractable" in d["status_detail"].lower()


def test_corrupt_docx_fails(env):
    p = env["files"] / "broken.docx"
    p.write_bytes(b"not a zip at all")
    d = _add(p)
    assert d["status"] == "failed"


def test_unsupported_type_rejected(env):
    with pytest.raises(ValueError):
        ingest.add_file("photo.jpg", b"\xff\xd8\xff", sync=True)


def test_scanned_page_is_reported_not_read(env):
    """A PDF whose page 2 is an image only: that page must be marked unreadable."""
    p = env["files"] / "scan.pdf"
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 545, 400),
                        "Results. Training to failure produced similar strength gains compared with non-failure training.",
                        fontsize=11)
    img_page = doc.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), 0)
    pix.set_rect(pix.irect, (120, 120, 120))
    img_page.insert_image(pymupdf.Rect(50, 50, 400, 400), pixmap=pix)
    doc.save(str(p))
    d = _add(p)
    rep = d["extraction_report"]
    assert rep["units"][1]["status"] == "unreadable"
    assert "NOT read" in rep["units"][1]["warnings"][0] or "OCR" in rep["units"][1]["warnings"][0]
    assert d["status"] == "needs_review" and d["pages_failed"] == 1 and d["pages_ok"] == 1


def test_two_column_reading_order(env):
    p = F.make_pdf(env["files"] / "cols.pdf", [[
        "Left column first paragraph about squats.", "Left column second paragraph about lunges.",
        "Right column first paragraph about rows.", "Right column second paragraph about presses."]],
        two_column_page=0)
    res = extract(p)
    text = res.units[0].text
    assert text.index("Left column second") < text.index("Right column first")


def test_docx_pptx_xlsx_csv_md_extraction(env):
    import docx
    import openpyxl
    from pptx import Presentation

    d = docx.Document()
    d.add_heading("Protein Intake", level=1)
    d.add_paragraph("Participants consumed 1.6 g/kg/day of protein during the 10-week program.")
    d.save(env["files"] / "nutrition.docx")
    prs = Presentation()
    s = prs.slides.add_slide(prs.slide_layouts[1])
    s.shapes.title.text = "Rest Intervals"
    s.placeholders[1].text = "Rest intervals of 2 minutes were used between sets."
    prs.save(env["files"] / "rest.pptx")
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["study", "weeks", "outcome"])
    ws.append(["Trial X", 8, "strength increased"])
    wb.save(env["files"] / "table.xlsx")
    (env["files"] / "log.csv").write_text("exercise,sets,reps\nSquat,3,5\n", encoding="utf-8")
    (env["files"] / "notes.md").write_text("# Mobility\n\nDaily stretching sessions lasted 10 minutes.\n", encoding="utf-8")
    for name in ("nutrition.docx", "rest.pptx", "table.xlsx", "log.csv", "notes.md"):
        doc = _add(env["files"] / name)
        assert doc["status"] == "processed", (name, doc["status_detail"])
    r = rag.ask("How long were rest intervals between sets?", lang="en")
    assert r["status"] == "answered"
    c = next(iter(r["citations"].values()))
    assert c["filename"] == "rest.pptx" and c["unit_index"] == 1 and c["location"] == "slide 1"
    r = rag.ask("protein intake per day", lang="en")
    assert any(c["filename"] == "nutrition.docx" and c["section"] == "Protein Intake" for c in r["citations"].values())


def test_replace_file_reprocesses_without_duplicates(env):
    p = F.study_a(env["files"])
    d = _add(p)
    new = env["files"] / "v2.txt"
    new.write_text("Results\n\nThe revised analysis found that 12 sets per week maximised quadriceps growth in trained men.\n",
                   encoding="utf-8")
    d2 = ingest.replace_file(d["id"], new.read_bytes(), filename="Volume_Study_A_v2.txt", sync=True)
    assert d2["id"] == d["id"] and len(ingest.list_documents()) == 1
    r = rag.ask("How many sets per week maximised quadriceps growth?", lang="en")
    assert "12 sets per week" in " ".join(c["text"] for c in r["claims"])
    assert "15 sets per week" not in " ".join(c["text"] for c in r["claims"])


def test_same_name_conflict_asks_then_replaces(env):
    p = F.study_a(env["files"])
    _add(p)
    with pytest.raises(ingest.NameConflict):
        ingest.add_file(p.name, F.study_b(env["files"]).read_bytes(), sync=True)
    d = ingest.add_file(p.name, F.study_b(env["files"]).read_bytes(), on_name_conflict="replace", sync=True)
    assert len(ingest.list_documents()) == 1 and "Untrained Women" in d["title"]


def test_folder_import_and_incremental_add(env):
    folder = env["files"] / "lib"
    folder.mkdir()
    F.study_a(folder)
    F.study_b(folder)
    (folder / "ignore.exe").write_bytes(b"MZ")
    rep = ingest.import_folder(str(folder), sync=True)
    assert len(rep["added"]) == 2 and rep["skipped"] == ["ignore.exe"]
    first = {d["id"]: d["processed_at"] for d in ingest.list_documents()}
    F.injection_doc(folder)
    rep2 = ingest.import_folder(str(folder), sync=True)
    assert len(rep2["duplicates"]) == 2 and len(rep2["added"]) == 1
    for d in ingest.list_documents():
        if d["id"] in first:
            assert d["processed_at"] == first[d["id"]], "existing files must not be reprocessed"


def test_rebuild_index_from_originals(env):
    _add(F.study_a(env["files"]))
    out = ingest.rebuild_index()
    assert out["documents"] == 1 and out["processed"] == 1
    assert rag.ask("How many sets per week did the high-volume group perform?", lang="en")["status"] == "answered"


def test_metadata_detection(env):
    d = _add(F.study_b(env["files"]))
    assert d["title"] == "Moderate Versus High Training Volume in Untrained Women"
    assert d["authors"] == "Smith K" and d["year"] == 2021
    a = _add(F.study_a(env["files"]))
    assert a["doc_type"] == "rct" and a["year"] is None  # not stated -> not guessed


# ------------------------------------------------------------------ verifier

def test_verifier_rules():
    passage = "Training volume was associated with greater muscle thickness in 34 trained men over 8 weeks."
    assert locate_quote("associated with greater muscle thickness", passage)
    assert locate_quote("this sentence is not in the passage at all", passage) is None
    ok, errors, _ = check_claim("Volume was associated with greater muscle thickness in 34 men.", [passage])
    assert ok
    ok, errors, _ = check_claim("Volume was associated with greater muscle thickness in 50 men.", [passage])
    assert not ok and "numbers" in errors[0]
    ok, errors, _ = check_claim("Higher volume causes greater muscle thickness.", [passage])
    assert not ok and "causal" in errors[0]


def test_llm_output_is_verified(env, fake_llm):
    _add(F.study_a(env["files"]))

    def responder(system, user, n):
        if n == 1:
            ids = re.findall(r'<passage id="(E\d+)"', user)
            res_id = next(i for i in ids if "15 sets per week" in user.split(f'id="{i}"')[1].split("</passage>")[0])
            return {"claims": [
                {"text": "The 15 sets per week group gained more quadriceps thickness than the 5 sets group.", "type": "stated",
                 "evidence": [{"id": res_id, "quote": "The high-volume group (15 sets per week) showed significantly greater increases in quadriceps muscle thickness"}]},
                {"text": "The 30 sets per week group gained the most.", "type": "stated",
                 "evidence": [{"id": res_id, "quote": "The high-volume group (15 sets per week) showed significantly greater increases"}]},
                {"text": "A study by Brown 2019 found the same.", "type": "stated", "evidence": [{"id": "E99", "quote": "Brown 2019"}]},
                {"text": "Women respond the same way.", "type": "inference",
                 "evidence": [{"id": res_id, "quote": "fabricated quote that does not exist in the passage"}]},
            ], "conflicts": [], "gaps": ["No data on women"]}
        return {"results": [{"index": 0, "verdict": "supported"}]}

    llm = fake_llm(responder)
    r = rag.ask("How many sets per week did the high-volume group perform?", lang="en", llm=llm)
    assert r["status"] == "answered" and r["mode"] == "llm:fake"
    assert [c["text"] for c in r["claims"]] == ["The 15 sets per week group gained more quadriceps thickness than the 5 sets group."]
    reasons = json.dumps(r["rejected_claims"])
    assert "numbers not found" in reasons and "unknown passage" in reasons and "quote not found" in reasons
    eid = r["claims"][0]["citations"][0]
    assert r["citations"][eid]["pdf_page"] == 3


def test_judge_rejects_unsupported(env, fake_llm):
    _add(F.study_a(env["files"]))

    def responder(system, user, n):
        if n == 1:
            sent = "Training volume showed a dose-response relationship with muscle hypertrophy"
            eid = next(i for i in re.findall(r'<passage id="(E\d+)"', user)
                       if sent in user.split(f'id="{i}"')[1].split("</passage>")[0].replace("\n", " "))
            return {"claims": [{"text": "تشير الدراسة إلى نتيجة ما", "type": "stated", "evidence": [{"id": eid, "quote": sent}]}]}
        return {"results": [{"index": 0, "verdict": "unsupported"}]}

    r = rag.ask("weekly sets hypertrophy trained men", lang="ar", llm=fake_llm(responder))
    assert r["status"] == "insufficient" and r["rejected_claims"]


# ------------------------------------------------------------------ privacy

def test_external_provider_requires_consent(env):
    update_settings({"llm_provider": "anthropic"})
    import os
    os.environ["ANTHROPIC_API_KEY"] = "test-key-not-real"
    try:
        with pytest.raises(PrivacyBlocked):
            get_llm(get_settings(), "library")
        update_settings({"allow_external_llm": True})
        assert get_llm(get_settings(), "library") is not None
        with pytest.raises(PrivacyBlocked):
            get_llm(get_settings(), "client")
    finally:
        del os.environ["ANTHROPIC_API_KEY"]
        update_settings({"llm_provider": "none", "allow_external_llm": False})


def test_settings_api_never_returns_secrets(env, monkeypatch):
    from fastapi.testclient import TestClient
    from navcoach.main import app
    monkeypatch.setenv("ANTHROPIC_API_KEY", "sk-secret-value-123")
    c = TestClient(app)
    s = c.get("/api/settings").json()
    assert "sk-secret-value-123" not in json.dumps(s) and s["secrets"]["ANTHROPIC_API_KEY"] == "configured"
    r = c.post("/api/settings", json={"llm_provider": "anthropic"})
    assert r.status_code == 400 and r.json()["detail"] == "external_model_requires_consent"
    assert get_settings().llm_provider == "none"


# ------------------------------------------------------------------ clients & programs

def test_complete_client_program_is_evidence_linked(env):
    _add(F.study_a(env["files"]))
    _add(F.study_b(env["files"]))
    c = clients.create_client("Omar")
    p = F.client_complete(env["files"])
    clients.add_client_file(c["id"], p.name, p.read_bytes())
    a = clients.analyze(c["id"], lang="en")
    assert a["readiness"] == "ready" and a["red_flags"] == []
    assert len(a["training_log"]) == 2 and a["training_log"][0]["effort_scale"] == "RIR"
    prog = programs.build_program(c["id"], lang="en")
    assert prog["id"] and prog["version"] == 1
    content = prog["content"]
    els = {e["key"]: e for e in content["elements"]}
    for e in content["elements"]:
        for v in e["evidence_values"]:
            cit = content["citations"][v["citation"]]
            assert v["value"] in v["quote"] and v["quote"].replace("\n", " ")[:40] in cit["passage"].replace("\n", " ")
        if e["basis"] == "evidence":
            assert e["citations"]
    assert els["frequency"]["basis"] == "evidence"
    assert els["frequency"]["proposed"].startswith("3 days")  # "three days per week" in source, within 3 available days
    assert els["frequency"]["evidence_values"][0]["value"] == "three days per week"
    assert els["weekly_volume"]["basis"] == "evidence"
    assert els["reps"]["basis"] == "evidence" and "8 to 12 repetitions" in els["reps"]["proposed"]
    vol_files = {content["citations"][v["citation"]]["filename"] for v in els["weekly_volume"]["evidence_values"]}
    assert vol_files == {"Volume_Study_A.pdf", "Volume_Study_B.pdf"}  # both sources, nothing hidden
    # Revision keeps history and records the reason
    p2 = programs.revise_program(prog["id"], [{"key": "rest", "proposed": "2 minutes", "coach_notes": "client preference"}],
                                 reason="Time constraint")
    assert p2["version"] == 2 and p2["change_notes"][0]["reason"] == "Time constraint"
    assert programs.get_program(prog["id"])["content"]["elements"] != p2["content"]["elements"]
    md = programs.program_markdown(p2, "en")
    assert "| Element |" in md and "Volume_Study_A.pdf" in md


def test_population_mismatch_is_flagged(env):
    _add(F.study_a(env["files"]))
    c = clients.create_client("Lina", profile={
        "goal": "muscle hypertrophy", "experience": "beginner, no experience", "days_available": "4",
        "equipment": "gym", "injuries": "none", "medical": "none", "sex": "female"})
    prog = programs.build_program(c["id"], lang="en")
    vol = next(e for e in prog["content"]["elements"] if e["key"] == "weekly_volume")
    notes = " ".join(vol["applicability"])
    assert "'trained' while the client is 'untrained'" in notes and "'men' while the client is 'women'" in notes


def test_red_flags_block_program_and_refer(env):
    c = clients.create_client("Ali", profile={
        "goal": "strength", "experience": "2 years", "days_available": "3", "equipment": "gym",
        "injuries": "sharp pain in the lower back when lifting", "medical": "none"})
    a = clients.analyze(c["id"], lang="en", evidence=False)
    assert any(f["flag"] == "acute_pain" for f in a["red_flags"])
    assert a["readiness"] == "needs_medical_clearance"
    res = programs.build_program(c["id"], lang="en")
    assert res["status"] == "needs_medical_clearance" and "id" not in res
    # negations are not flagged
    c2 = clients.create_client("Mona", profile={"injuries": "none", "medical": "no chest pain, no dizziness"})
    assert clients.analyze(c2["id"], evidence=False)["red_flags"] == []


def test_client_data_is_isolated(env):
    c1 = clients.create_client("A")
    c2 = clients.create_client("B")
    p = F.client_complete(env["files"])
    clients.add_client_file(c1["id"], p.name, p.read_bytes())
    a2 = clients.analyze(c2["id"], evidence=False)
    assert a2["facts"] == [] and a2["training_log"] == []
    assert (env["data"] / "clients" / c1["id"]).exists()
    assert not any((env["data"] / "clients" / c2["id"]).iterdir())


# ------------------------------------------------------------------ API, backup

def test_api_end_to_end(env):
    from fastapi.testclient import TestClient
    from navcoach.main import app
    c = TestClient(app)
    assert "Nav Coaching" in c.get("/").text
    p = F.study_a(env["files"])
    r = c.post("/api/documents", files=[("files", (p.name, p.read_bytes(), "application/pdf"))])
    assert r.json()["results"][0]["result"] == "queued"
    ingest.worker.join()
    docs = c.get("/api/documents").json()
    assert docs[0]["status"] == "processed"
    r = c.post("/api/documents", files=[("files", (p.name, p.read_bytes(), "application/pdf"))])
    assert r.json()["results"][0]["result"] == "duplicate"
    ans = c.post("/api/ask", json={"question": "How many sets per week did the high-volume group perform?", "lang": "en"}).json()
    assert ans["status"] == "answered"
    cid = next(iter(ans["citations"].values()))["chunk_id"]
    assert c.get(f"/api/chunks/{cid}").json()["chunk"]["id"] == cid
    assert c.get(f"/api/documents/{docs[0]['id']}/file").status_code == 200
    assert c.get("/api/history").json()[0]["question"].startswith("How many")
    assert c.delete(f"/api/documents/{docs[0]['id']}").status_code == 400  # confirmation required
    assert c.delete(f"/api/documents/{docs[0]['id']}?confirm=true").json()["deleted"]
    assert c.get(f"/api/chunks/{cid}").status_code == 404
    cmp_ = c.post("/api/compare", json={"doc_ids": [docs[0]["id"]]})
    assert cmp_.status_code == 400


def test_compare_table_uses_only_file_content(env):
    a = _add(F.study_a(env["files"]))
    b = _add(F.study_b(env["files"]))
    from navcoach.compare import compare
    r = compare([a["id"], b["id"]], "training volume hypertrophy", "en")
    rows = {row["filename"]: row for row in r["rows"]}
    A, B = rows["Volume_Study_A.pdf"]["cells"], rows["Volume_Study_B.pdf"]["cells"]
    assert ("34" in A["sample"]["text"] or "Thirty-four" in A["sample"]["text"]) and "8" in A["duration"]["text"]
    assert "ultrasound" in A["measures"]["text"]
    assert "Limitations include" in A["limitations"]["text"]
    assert B["measures"]["text"] == "Not stated in the file"  # not invented
    for row in r["rows"]:
        for cell in row["cells"].values():
            if cell["citation"] and cell["citation"] in r["citations"] and r["citations"][cell["citation"]]["quote"]:
                assert cell["text"] == r["citations"][cell["citation"]]["quote"]


def test_backup_and_restore(env):
    from navcoach import backup
    _add(F.study_a(env["files"]))
    cl = clients.create_client("Keep")
    path = backup.create_backup()
    z = zipfile.ZipFile(path)
    assert "navcoach.db" in z.namelist() and any(n.startswith("library/") for n in z.namelist())
    assert not any(n.endswith(".env") for n in z.namelist())
    data = path.read_bytes()
    docs = ingest.list_documents()
    ingest.delete_document(docs[0]["id"])
    clients.delete_client(cl["id"])
    assert ingest.list_documents() == []
    backup.restore_backup(data)
    assert len(ingest.list_documents()) == 1 and clients.list_clients()[0]["name"] == "Keep"
    assert rag.ask("How many sets per week did the high-volume group perform?", lang="en")["status"] == "answered"
    with pytest.raises(ValueError):
        backup.restore_backup(b"not a zip")
