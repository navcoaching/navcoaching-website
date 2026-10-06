"""Summaries & Slides / Study mode — acceptance tests (FILE-ONLY knowledge policy)."""
import json
import os
import re
import subprocess
import sys
import time
from pathlib import Path

import pymupdf
import pytest

from navcoach import db, docmodel, ingest, study, summarize
from navcoach.textutil import fold
from navcoach.verify import locate_quote
from tests import fixtures as F

ROOT = Path(__file__).resolve().parents[1]
NOT_FOUND_AR = "لم أجد هذه المعلومة في الملفات المتاحة لدي."


def _add(path):
    return ingest.add_file(path.name, path.read_bytes(), sync=True)


def _summary(doc, level="detailed", lang="en", **kw):
    s = summarize.create(doc["id"], level, lang, sync=True, **kw)
    assert s["status"] == "done", s.get("error")
    return summarize.get(s["id"], include_answers=True)


def _all_sourced_items(s):
    """Every (text, source) pair shown anywhere: summary, slides, notes, study questions."""
    out = []
    for c in summarize.iter_claims(s["content"]):
        for src in c["sources"]:
            out.append((c["text"], c.get("kind"), src))
    for sl in s["slides"]["slides"]:
        for b in sl["bullets"] + sl["notes"]:
            for src in b.get("sources", []):
                out.append((b.get("full") or b["text"], b.get("kind"), src))
    for q in s["study"]["questions"]:
        out.append((q["explanation"], "question", q["source"]))
    return out


def _doc_text(doc_id):
    with db.session() as conn:
        return " ".join(r["text"] for r in conn.execute("SELECT text FROM chunks WHERE doc_id=?", (doc_id,)))


def _long_pdf(path: Path, sections: int = 24) -> Path:
    pages = []
    for k in range(1, sections + 1):
        pages.append([f"{k}. Training Topic {k}",
                      f"This section examines training variable {k} in resistance-trained adults.",
                      f"Participants in protocol {k} completed {k + 2} sets per week for {k + 4} weeks.",
                      f"The results showed that variable {k} increased strength by {k * 1.5:.1f}% compared with baseline.",
                      f"Limitations of protocol {k} include the small sample of {k + 10} participants."] )
    return F.make_pdf(path, pages, title="Long Training Handbook")


# 1 ------------------------------------------------------------------------- short research file
def test_1_short_research_file_full_summary(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc)
    c = s["content"]
    assert [q["key"] for q in c["quick"]] == ["topic", "aim", "main_finding", "why_important"]
    assert all(q["sources"] for q in c["quick"])
    assert any("8.1%" in x["text"] for x in c["findings"])
    assert any(x["term"].startswith("Repetitions in reserve") for x in c["concepts"])
    assert "ultrasound" in c["methodology"]["measurement_method"]["text"]
    assert "n = 24" in c["methodology"]["participants"]["text"]
    assert any("small sample" in x["text"] for x in c["limitations"])
    assert any("Coaches can prescribe" in x["text"] for x in c["practical"])
    assert c["misconceptions"] and "commonly believed" in c["misconceptions"][0]["text"]
    assert s["verification"]["rejected"] == [] and s["verification"]["checked"] > 20
    ov = c["overview"]
    assert ov["authors"] == "Lee M, Park S" and ov["pages"] == 4 and ov["title"].startswith("Training to Failure")
    # key points are ranked by importance, not by position
    imps = [x["importance"] for x in c["key_points"]]
    assert imps == sorted(imps, reverse=True)
    assert [x["sources"][0]["page"] for x in c["key_points"]] != sorted(x["sources"][0]["page"] for x in c["key_points"])


# 2 ------------------------------------------------------------------------- long file (hierarchical)
def test_2_long_file_is_processed_hierarchically(env, fake_llm):
    doc = _add(_long_pdf(env["files"] / "Long_Handbook.pdf"))
    t0 = time.time()
    s = _summary(doc, "expert")
    assert time.time() - t0 < 60
    assert len(s["content"]["outline"]) >= 20
    assert len(s["content"]["section_summaries"]) >= 20
    sizes = []

    def responder(system, user, n):
        if "VERIFIED POINTS" in user:
            return {"topic": None, "aim": None, "main_finding": None, "why_important": None}
        sizes.append(len(user))
        pid, body = re.findall(r'<passage id="(P\d+)"[^>]*>\n(.*?)\n</passage>', user, re.S)[0]
        sent = re.search(r"The results showed that variable \d+ increased strength by [\d.]+% compared with baseline\.", body)
        return {"points": [{"text": sent.group(0), "category": "finding", "importance": 4,
                            "evidence": [{"id": pid, "quote": sent.group(0)}]}] if sent else []}

    llm = fake_llm(responder)
    s2 = summarize.create(doc["id"], "standard", "en", sync=True, llm=llm)
    s2 = summarize.get(s2["id"])
    assert len(sizes) > 1, "a long document must be summarised in several sections (map step)"
    assert max(sizes) < 9000, "no single request may contain the whole document"
    assert s2["mode"] == "llm:fake" and s2["content"]["findings"]


# 3 ------------------------------------------------------------------------- tables
def test_3_tables_become_charts_with_file_values_only(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc)
    t = s["content"]["tables"][0]
    assert t["header"] == ["Group", "CSA change %", "Fatigue score"] and t["numeric_columns"] == [1, 2]
    chart = next(sl for sl in s["slides"]["slides"] if sl["type"] == "chart")
    v = chart["visual"]
    assert v["labels"] == ["Failure", "2 RIR", "Control"]
    assert v["series"][0]["values"] == [8.1, 7.4, 0.6] and v["series"][1]["values"] == [7.2, 5.1, 2.0]
    chunk_text = _doc_text(doc["id"])
    for series in v["series"]:
        for val in series["values"]:
            assert str(val) in chunk_text
    assert v["source"]["page"] == 3


def test_3b_no_chart_without_table_data(env):
    doc = _add(F.study_a(env["files"]))
    s = _summary(doc)
    assert not any(sl["type"] == "chart" for sl in s["slides"]["slides"])


# 4 ------------------------------------------------------------------------- multiple sections
def test_4_outline_and_sections(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc)
    outline = {o["title"]: o for o in s["content"]["outline"]}
    assert outline["2.1 Participants"]["level"] == 2 and outline["3. Results"]["level"] == 1
    assert outline["3. Results"]["page_start"] == 3 and outline["3. Results"]["first_chunk_id"]
    titles = [x["title"] for x in s["content"]["section_summaries"]]
    assert "3. Results" in titles and "Limitations" in titles
    ex = study.explain_claims(docmodel.load(doc["id"]), *study.section_claims(s["content"], "findings")[::1], lang="en")
    parts = {p["key"]: p for p in ex["parts"]}
    assert parts["meaning"]["items"] and parts["finding"]["items"] and parts["takeaway"]["items"]


# 5 ------------------------------------------------------------------------- insufficient information
def test_5_missing_information_is_declared(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc, lang="ar")
    assert s["content"]["methodology"]["design"]["kind"] == "not_found"
    assert s["content"]["methodology"]["design"]["text"] == NOT_FOUND_AR
    r = study.ask_document(doc["id"], doc["filename"], "What is the effect of caffeine on sprint performance?", "ar")
    assert r["status"] == "insufficient" and r["message"] == NOT_FOUND_AR and r["claims"] == []
    m = docmodel.load(doc["id"])
    out = study.slide_chat(m, s["slides"], "sl2", "أعطني مثالًا على هذه الفكرة", "ar")
    ex = out["explanation"]["parts"][0]
    assert ex["key"] == "example" and not ex["items"] and "مثال" in ex["message"]
    doc2 = _add(F.study_a(env["files"]))
    s2 = _summary(doc2, lang="ar")
    assert s2["content"]["limitations"] or s2["content"]["limitations_message"]


# 6 ------------------------------------------------------------------------- two conflicting files
def test_6_conflicting_files_stay_separate(env):
    a = _add(F.study_a(env["files"]))
    b = _add(F.study_b(env["files"]))
    sa = _summary(a)
    assert {src["doc_id"] for _, _, src in _all_sourced_items(sa)} == {a["id"]}
    r = study.ask_document(a["id"], a["filename"], "Did increasing volume beyond 10 weekly sets increase hypertrophy in untrained women?", "en")
    assert all(c["filename"] == "Volume_Study_A.pdf" for c in r["citations"].values())
    from navcoach import rag
    lib = rag.ask("ماذا تقول ملفاتي عن أثر زيادة الحجم التدريبي على التضخم العضلي؟", lang="ar")
    assert lib["conflicts"]  # library-wide answers still expose the disagreement


# 7 ------------------------------------------------------------------------- scanned pages
def test_7_scanned_pages_are_excluded_and_reported(env):
    p = env["files"] / "scan_mixed.pdf"
    d = pymupdf.open()
    pg = d.new_page()
    pg.insert_textbox(pymupdf.Rect(50, 50, 545, 400), "Results\n\nThe results showed that sprint time decreased by 3% after 6 weeks of training.", fontsize=11)
    img = d.new_page()
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, 200, 200), 0)
    pix.set_rect(pix.irect, (90, 90, 90))
    img.insert_image(pymupdf.Rect(50, 50, 400, 400), pixmap=pix)
    d.save(str(p))
    doc = _add(p)
    s = _summary(doc, lang="en")
    assert s["content"]["overview"]["pages_unread"] == [2]
    assert any("unread pages: 2" in w for w in s["content"]["warnings"])
    assert all(src["page"] != 2 for _, _, src in _all_sourced_items(s))


# 8 ------------------------------------------------------------------------- every source is correct
def test_8_every_claim_source_is_real_and_on_the_right_page(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc, "expert")
    pdf = pymupdf.open(env["data"] / ingest.get_document(doc["id"])["stored_path"])
    items = _all_sourced_items(s)
    assert len(items) > 40
    with db.session() as conn:
        for text, kind, src in items:
            row = conn.execute("SELECT * FROM chunks WHERE id=?", (src["chunk_id"],)).fetchone()
            assert row is not None and row["doc_id"] == doc["id"]
            assert row["page"] == src["page"]
            quote = src["quote"]
            assert locate_quote(quote, row["text"]), quote
            page_text = fold(pdf[src["page"] - 1].get_text())
            assert fold(quote)[:50] in page_text
            if kind == "quote":
                assert locate_quote(text, row["text"]), text


# 9 ------------------------------------------------------------------------- slides contain only file content
def test_9_slides_contain_only_source_content(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc)
    text = _doc_text(doc["id"])
    doc_numbers = set(re.findall(r"\d+(?:\.\d+)?", text))
    for sl in s["slides"]["slides"]:
        for b in sl["bullets"]:
            if b.get("kind") == "not_found":
                continue
            full = b.get("full") or b["text"]
            if b.get("sources"):
                assert any(locate_quote(full, ingest_chunk(b2["chunk_id"])) for b2 in b["sources"]), full
            assert set(re.findall(r"\d+(?:\.\d+)?", full)) <= doc_numbers
            if b.get("highlight"):
                assert b["highlight"] in full
        for n in sl["notes"]:
            if n.get("sources"):
                assert any(locate_quote(n["text"], ingest_chunk(x["chunk_id"])) for x in n["sources"])
        v = sl.get("visual") or {}
        for node in v.get("nodes") or []:
            if isinstance(node, dict):
                assert node["sources"]
        for e in v.get("edges") or []:
            assert e["sources"] and e["label"].lower() in fold(" ".join(x["quote"] for x in e["sources"]))


def ingest_chunk(chunk_id):
    with db.session() as conn:
        return conn.execute("SELECT text FROM chunks WHERE id=?", (chunk_id,)).fetchone()["text"]


# 10 ------------------------------------------------------------------------ study questions only from file
def test_10_study_questions_come_from_the_file(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc, "expert", lang="ar")
    qs = s["study"]["questions"]
    assert {q["type"] for q in qs} == {"mcq", "tf", "short"}
    text = _doc_text(doc["id"])
    nums = set(re.findall(r"\d+(?:\.\d+)?", text))
    for q in qs:
        assert locate_quote(q["source"]["quote"], ingest_chunk(q["source"]["chunk_id"]))
        if q["type"] == "mcq":
            for o in q["options"]:
                assert o in nums or o.lower() in text.lower(), o
        if q["type"] == "tf":
            assert set(re.findall(r"\d+(?:\.\d+)?", q["statement"])) <= nums
            if q["answer"] is True:
                assert locate_quote(q["statement"], ingest_chunk(q["source"]["chunk_id"]))
    public = summarize.get(s["id"])["study"]["questions"]
    assert all("answer" not in q and "source" not in q for q in public)
    mcq = next(q for q in qs if q["type"] == "mcq")
    assert study.check_answer(s["study"], mcq["id"], mcq["answer"])["correct"] is True
    assert study.check_answer(s["study"], mcq["id"], (mcq["answer"] + 1) % len(mcq["options"]))["correct"] is False
    sa = next(q for q in qs if q["type"] == "short" and q["accept"].get("numbers"))
    assert study.check_answer(s["study"], sa["id"], f"تقريبًا {sa['accept']['numbers'][0]}")["correct"] is True
    tf = next(q for q in qs if q["type"] == "tf")
    assert study.check_answer(s["study"], tf["id"], tf["answer"])["correct"] is True


# 11 ------------------------------------------------------------------------ persistence across restart
def test_11_summary_and_slides_persist_after_restart(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc)
    code = ("import json, sys; sys.path.insert(0, %r)\nfrom navcoach import summarize\n"
            "s = summarize.get(%r)\nprint(json.dumps({'status': s['status'], 'slides': len(s['slides']['slides']),"
            " 'kp': len(s['content']['key_points']), 'q': len(s['study']['questions'])}))\n" % (str(ROOT), s["id"]))
    out = subprocess.run([sys.executable, "-c", code], env={**os.environ, "NAV_DATA_DIR": str(env["data"])},
                         capture_output=True, text=True, timeout=60)
    assert out.returncode == 0, out.stderr
    res = json.loads(out.stdout.strip().splitlines()[-1])
    assert res == {"status": "done", "slides": len(s["slides"]["slides"]), "kp": len(s["content"]["key_points"]),
                   "q": len(s["study"]["questions"])}


# 12 ------------------------------------------------------------------------ replaced / deleted source file
def test_12_replaced_file_marks_summary_stale_and_regenerate_updates(env):
    doc = _add(F.rich_study(env["files"]))
    s = _summary(doc)
    assert summarize.get(s["id"])["stale"] is False
    new = env["files"] / "v2.txt"
    new.write_text("Results\n\nThe revised analysis found that 3 sets per week produced 5.5% greater hypertrophy in trained women.\n",
                   encoding="utf-8")
    ingest.replace_file(doc["id"], new.read_bytes(), filename="Failure_v2.txt", sync=True)
    old = summarize.get(s["id"])
    assert old["stale"] is True
    assert summarize.list_all(doc["id"])[0]["stale"] is True
    summarize.regenerate(s["id"], sync=True)
    fresh = summarize.get(s["id"], include_answers=True)
    assert fresh["stale"] is False
    assert any("5.5%" in x["text"] for x in fresh["content"]["findings"])
    assert not any("8.1%" in t for t, _, _ in _all_sourced_items(fresh))
    ingest.delete_document(doc["id"])
    assert summarize.get(s["id"]) is None  # summaries are deleted with their file


# ------------------------------------------------------------------------- model path is verified
def test_model_summary_rejects_unsupported_claims(env, fake_llm):
    doc = _add(F.rich_study(env["files"]))

    def responder(system, user, n):
        if "VERIFIED POINTS" in user:
            return {"topic": {"text": "The study compared failure training with RIR training in women.", "points": ["C1"]},
                    "aim": {"text": "It proved creatine doubles strength in 2 weeks.", "points": ["C1"]},
                    "main_finding": None, "why_important": None}
        points = []
        for pid, body in re.findall(r'<passage id="(P\d+)"[^>]*>\n(.*?)\n</passage>', user, re.S):
            m = re.search(r"The failure group increased cross-sectional area by 8\.1%[^.]*\.", body.replace("\n", " "))
            if m:
                points += [
                    {"text": "Failure training increased cross-sectional area by 8.1%, similar to 2 RIR (7.4%).",
                     "category": "finding", "importance": 5, "evidence": [{"id": pid, "quote": m.group(0)}]},
                    {"text": "Failure training increased cross-sectional area by 12%.", "category": "finding",
                     "importance": 5, "evidence": [{"id": pid, "quote": m.group(0)}]},
                    {"text": "Training to failure causes injuries.", "category": "key", "importance": 5,
                     "evidence": [{"id": pid, "quote": "Training to failure causes injuries in most athletes"}]},
                ]
        return {"points": points}

    s = summarize.create(doc["id"], "standard", "en", sync=True, llm=fake_llm(responder))
    s = summarize.get(s["id"])
    texts = [c["text"] for c in summarize.iter_claims(s["content"])]
    assert "Failure training increased cross-sectional area by 8.1%, similar to 2 RIR (7.4%)." in texts
    assert not any("12%" in t or "causes injuries" in t or "creatine" in t for t in texts)
    reasons = json.dumps(s["verification"]["rejected"])
    assert "numbers not found" in reasons and "no verifiable quote" in reasons and "not supported" in reasons


def test_injected_instructions_never_reach_summary(env):
    doc = _add(F.injection_doc(env["files"]))
    s = _summary(doc)
    blob = json.dumps([t for t, _, _ in _all_sourced_items(s)]).lower()
    assert "ignore all previous" not in blob and "creatine" not in blob


# ------------------------------------------------------------------------- API & export
def test_summary_api_end_to_end(env):
    from fastapi.testclient import TestClient
    from navcoach.main import app
    c = TestClient(app)
    doc = _add(F.rich_study(env["files"]))
    r = c.post("/api/summaries", json={"doc_id": doc["id"], "level": "standard", "lang": "en"}).json()
    sid = r["id"]
    for _ in range(100):
        s = c.get(f"/api/summaries/{sid}").json()
        if s["status"] != "running":
            break
        time.sleep(0.1)
    assert s["status"] == "done"
    assert all("answer" not in q for q in s["study"]["questions"])
    ex = c.post(f"/api/summaries/{sid}/explain", json={"section": "findings", "lang": "en"}).json()
    assert [p["key"] for p in ex["parts"]] == ["meaning", "why_important", "finding", "takeaway"]
    q = s["study"]["questions"][0]
    assert "correct" in c.post(f"/api/summaries/{sid}/answer", json={"question_id": q["id"], "answer": 0}).json()
    chat = c.post(f"/api/summaries/{sid}/slide-chat", json={"slide_id": "sl2", "question": "What should I remember from this slide?"}).json()
    assert chat["kind"] == "explanation" and chat["explanation"]["parts"][0]["key"] == "takeaway"
    ask = c.post(f"/api/documents/{doc['id']}/ask", json={"question": "How long was the intervention?", "lang": "en"}).json()
    assert ask["status"] == "answered" and ask["scope"]["doc_id"] == doc["id"]
    pptx = c.get(f"/api/summaries/{sid}/export?format=pptx")
    assert pptx.status_code == 200 and pptx.content[:2] == b"PK"
    from pptx import Presentation
    import io
    prs = Presentation(io.BytesIO(pptx.content))
    assert len(prs.slides) == len(s["slides"]["slides"])
    assert "Source →" in prs.slides[5].notes_slide.notes_text_frame.text
    assert "## Key points" in c.get(f"/api/summaries/{sid}/export?format=md").text
    assert c.patch(f"/api/summaries/{sid}", json={"title": "My review"}).json()["title"] == "My review"
    assert c.delete(f"/api/summaries/{sid}").status_code == 400
    assert c.delete(f"/api/summaries/{sid}?confirm=true").json()["deleted"]
    assert c.get(f"/api/summaries/{sid}").status_code == 404


def test_concept_cards_are_real_terms_not_fragments(env):
    """Reported: concept cards like 'What', 'fat diet', 'Ketonix', 'consume medium-chain triglyceride (MCT)'."""
    doc = _add(F.keto_notes(env["files"]))
    s = _summary(doc, "expert", lang="ar")
    terms = [c["term"] for c in s["content"]["concepts"]]
    assert "Nutritional ketosis" in terms and "ketogenic diet" in terms
    assert "medium-chain triglyceride (MCT)" in terms
    for bad in ("What", "Ketonix", "A ketogenic diet", "A high fat diet"):
        assert bad not in terms
    assert not any(t.lower().startswith(("consume", "athletes", "what", "a ", "the ")) for t in terms)
    assert not any(len(t.split()) > 6 for t in terms)


def test_search_answers_stay_on_topic(env):
    """Reported: answers unrelated to the question. Generic words must not carry a match."""
    from navcoach import rag
    _add(F.keto_notes(env["files"]))
    _add(F.study_a(env["files"]))
    _add(F.rich_study(env["files"]))
    r = rag.ask("كم عدد أيام تدريب المقاومة؟", lang="ar")
    texts = [c["text"] for c in r["claims"]]
    assert texts and all(re.search(r"days per week|times per week", t) for t in texts), texts
    r = rag.ask("What happens to sprint performance on a ketogenic diet?", lang="en")
    assert [c["text"] for c in r["claims"]] == [
        "Sprint performance decreased by 4% during the ketogenic diet compared with a high carbohydrate diet."]
    r = rag.ask("ما تأثير الكيتو على الأداء؟", lang="ar")
    assert r["claims"] and all("ketogenic" in c["text"] for c in r["claims"])
