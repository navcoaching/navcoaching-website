"""Format-specific extractors."""
from __future__ import annotations

import csv
import io
import re
import shutil
import statistics
from pathlib import Path

from ..textutil import clean_text
from . import ExtractionError, ExtractResult, Unit

SECTION_WORDS = (
    "abstract", "introduction", "background", "methods", "methodology", "materials and methods", "participants",
    "results", "discussion", "conclusion", "conclusions", "limitations", "practical applications",
    "references", "acknowledgements", "keywords", "summary", "training volume", "statistical analysis",
    "الملخص", "المقدمة", "الطرق", "المنهجية", "النتائج", "المناقشة", "الخلاصة", "القيود", "المراجع",
)
_SECTION_RE = re.compile(r"^\s*(\d+(\.\d+)*\.?\s+)?(" + "|".join(re.escape(w) for w in SECTION_WORDS) + r")\s*:?\s*$", re.I)


def _looks_like_heading(line: str) -> bool:
    return bool(_SECTION_RE.match(line))


# --------------------------------------------------------------------------- PDF

def _ocr_available() -> bool:
    return shutil.which("tesseract") is not None


_WORDLIKE = re.compile(r"[^\W\d_]{2,}")
_NUM_OR_PUNCT = re.compile(r"[\d\W_]+")


def _garbled(text: str) -> bool:
    """True when the text layer is unusable (broken font encoding).

    Pages made mostly of numbers, dot leaders or table cells (contents pages,
    data tables) are NOT garbled: only tokens that contain something other than
    digits/punctuation are judged, and those must mostly look like words.
    """
    if not text:
        return False
    bad = sum(1 for c in text if c == "\ufffd" or "\ue000" <= c <= "\uf8ff" or (ord(c) < 32 and c not in "\n\t"))
    if bad / len(text) > 0.05:
        return True
    # Symbol soup (no letters or digits): ignore dot leaders / rules first.
    core = re.sub(r"[.\u00b7\u2026_\-\u2013\u2014=|]{3,}|\s+", "", text)
    if len(core) > 200 and sum(c.isalnum() for c in core) / len(core) < 0.4:
        return True
    judged = [t for t in text.split() if not _NUM_OR_PUNCT.fullmatch(t)]
    if len(judged) < 20:
        return False
    wordlike = sum(1 for t in judged if _WORDLIKE.search(t))
    return wordlike / len(judged) < 0.5


def _order_blocks(blocks: list[dict], page_width: float) -> list[dict]:
    """Reading order that respects two-column layouts.

    Full-width blocks act as separators; between separators, the left column is
    read fully before the right column.
    """
    mid = page_width / 2
    text_blocks = [b for b in blocks if b.get("type") == 0]
    text_blocks.sort(key=lambda b: (b["bbox"][1], b["bbox"][0]))
    ordered, left, right = [], [], []

    def flush():
        ordered.extend(sorted(left, key=lambda b: b["bbox"][1]))
        ordered.extend(sorted(right, key=lambda b: b["bbox"][1]))
        left.clear()
        right.clear()

    for b in text_blocks:
        x0, _, x1, _ = b["bbox"]
        if x1 <= mid + 10:
            left.append(b)
        elif x0 >= mid - 10:
            right.append(b)
        else:  # spans both columns
            flush()
            ordered.append(b)
    flush()
    return ordered


def _page_text_and_headings(page) -> tuple[str, list[tuple[int, str]], list[float]]:
    d = page.get_text("dict", flags=0)
    blocks = _order_blocks(d.get("blocks", []), page.rect.width)
    sizes = []
    for b in blocks:
        for ln in b.get("lines", []):
            for sp in ln.get("spans", []):
                if sp.get("text", "").strip():
                    sizes.extend([sp["size"]] * len(sp["text"]))
    body = statistics.median(sizes) if sizes else 10.0
    parts: list[str] = []
    headings: list[tuple[int, str]] = []
    offset = 0
    for b in blocks:
        lines = []
        for ln in b.get("lines", []):
            spans = ln.get("spans", [])
            t = "".join(sp.get("text", "") for sp in spans).strip()
            if not t:
                continue
            size = max((sp["size"] for sp in spans), default=body)
            bold = all(("bold" in sp.get("font", "").lower()) or (sp.get("flags", 0) & 16) for sp in spans)
            is_heading = (_looks_like_heading(t) or ((size >= body * 1.15 or bold) and len(t) < 90
                          and not t.endswith(".") and len(t.split()) <= 12 and not re.fullmatch(r"[\d\s.]+", t)))
            if is_heading:
                headings.append((offset + sum(len(x) + 1 for x in lines), t))
            lines.append(t)
        if lines:
            block_text = "\n".join(lines)
            parts.append(block_text)
            offset += len(block_text) + 2
    return "\n\n".join(parts), headings, sizes


_PRINTED_RE = re.compile(r"^(?:page\s+|p\.\s*|صفحة\s+)?(\d{1,4})$", re.I)


def _candidate_printed(text: str) -> str | None:
    lines = [ln.strip() for ln in text.split("\n") if ln.strip()]
    if not lines:
        return None
    for ln in lines[:2] + lines[-2:]:
        m = _PRINTED_RE.match(ln)
        if m:
            return m.group(1)
        m = re.match(r"^(\d{1,4})\s*[|•·]\s|\s[|•·]\s*(\d{1,4})$", ln)
        if m:
            return m.group(1) or m.group(2)
    return None


def extract_pdf(path: Path) -> ExtractResult:
    import pymupdf

    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise ExtractionError(f"PDF could not be opened ({type(exc).__name__})") from exc
    if doc.needs_pass:
        raise ExtractionError("PDF is password protected")
    res = ExtractResult(units=[], file_type="pdf")
    meta = doc.metadata or {}
    res.title = (meta.get("title") or "").strip() or None
    res.authors = (meta.get("author") or "").strip() or None
    m = re.match(r"D:(\d{4})", meta.get("creationDate") or "")
    meta_year = int(m.group(1)) if m else None
    has_labels = False
    try:
        has_labels = bool(doc.get_page_labels())
    except Exception:
        pass
    ocr = _ocr_available()
    candidates: list[tuple[int, int]] = []  # (index, printed number) from header/footer detection
    largest_first_page_line = None
    for i, page in enumerate(doc):
        unit = Unit(page=i + 1, text="")
        try:
            text, headings, sizes = _page_text_and_headings(page)
        except Exception:
            text, headings, sizes = "", [], []
            unit.warnings.append("text layer could not be parsed")
        text = clean_text(text)
        # Tables: append a structured rendering so cell values stay aligned.
        try:
            tabs = page.find_tables()
            for t_i, tab in enumerate(tabs.tables, start=1):
                rows = tab.extract()
                if not rows or len(rows) < 2:
                    continue
                rendered = "\n".join(" | ".join((c or "").replace("\n", " ").strip() for c in r) for r in rows)
                text += f"\n\n[Table {t_i} on this page]\n{rendered}"
        except Exception:
            unit.warnings.append("table detection failed")
        if i == 0 and sizes:
            largest_first_page_line = _largest_line(page)
        if len(text.strip()) < 25:
            has_images = bool(page.get_images(full=False))
            if ocr and has_images:
                try:
                    tp = page.get_textpage_ocr(full=True, dpi=300)
                    ocr_text = clean_text(page.get_text(textpage=tp))
                except Exception:
                    ocr_text = ""
                if len(ocr_text) >= 25:
                    unit.text, unit.status = ocr_text, "ocr"
                    unit.warnings.append("text obtained by OCR; verify against the original page")
                else:
                    unit.status = "unreadable"
                    unit.warnings.append("scanned page: OCR produced no reliable text")
            elif has_images:
                unit.status = "unreadable"
                unit.warnings.append("scanned/image-only page and OCR (tesseract) is not installed — page NOT read")
            else:
                unit.status = "empty"
                unit.warnings.append("blank page")
        elif _garbled(text):
            unit.status = "unreadable"
            unit.text = ""
            unit.warnings.append("text layer is garbled (bad font encoding) — page NOT read")
        else:
            unit.text = text
            unit.headings = headings
        if has_labels:
            try:
                label = page.get_label()
                if label:
                    unit.printed_page = label
            except Exception:
                pass
        if not unit.printed_page and unit.text:
            c = _candidate_printed(unit.text)
            if c and c.isdigit():
                candidates.append((i, int(c)))
        res.units.append(unit)
    # Accept header/footer page numbers only when a consistent offset is found.
    if candidates:
        offsets: dict[int, int] = {}
        for idx, num in candidates:
            offsets[num - idx] = offsets.get(num - idx, 0) + 1
        best, count = max(offsets.items(), key=lambda kv: kv[1])
        if count >= 2 or len(res.units) == 1:
            for idx, num in candidates:
                if num - idx == best:
                    res.units[idx].printed_page = str(num)
    if not res.title and largest_first_page_line:
        res.title = largest_first_page_line
    res._meta_year = meta_year  # type: ignore[attr-defined]
    doc.close()
    return res


def _largest_line(page) -> str | None:
    best, best_size = None, 0.0
    d = page.get_text("dict", flags=0)
    for b in d.get("blocks", []):
        for ln in b.get("lines", []):
            t = "".join(sp.get("text", "") for sp in ln.get("spans", [])).strip()
            size = max((sp["size"] for sp in ln.get("spans", [])), default=0)
            if t and len(t) > 8 and size > best_size:
                best, best_size = t, size
    return best


# -------------------------------------------------------------------------- DOCX

def extract_docx(path: Path) -> ExtractResult:
    import docx

    try:
        d = docx.Document(str(path))
    except Exception as exc:
        raise ExtractionError(f"DOCX could not be opened ({type(exc).__name__})") from exc
    res = ExtractResult(units=[], file_type="docx")
    cp = d.core_properties
    res.title = (cp.title or "").strip() or None
    res.authors = (cp.author or "").strip() or None
    if cp.created:
        res._meta_year = cp.created.year  # type: ignore[attr-defined]
    # Split into sections at headings; each section becomes a unit with paragraph locations.
    current_heading, buf, start_para = None, [], 1
    para_no = 0

    def flush(end_para):
        text = clean_text("\n\n".join(buf))
        if text:
            loc = f"paragraphs {start_para}-{end_para}" if end_para > start_para else f"paragraph {start_para}"
            u = Unit(page=None, text=text, location=loc)
            if current_heading:
                u.headings = [(0, current_heading)]
            res.units.append(u)

    for p in d.paragraphs:
        t = p.text.strip()
        if not t:
            continue
        para_no += 1
        style = (p.style.name or "").lower() if p.style is not None else ""
        if style.startswith("heading") or style == "title" or _looks_like_heading(t):
            flush(para_no - 1)
            buf, start_para, current_heading = [], para_no, t
            if style == "title" and not res.title:
                res.title = t
            continue
        buf.append(t)
    flush(para_no)
    for t_i, table in enumerate(d.tables, start=1):
        rows = [" | ".join(c.text.strip() for c in r.cells) for r in table.rows]
        if rows:
            res.units.append(Unit(page=None, text=f"[Table {t_i}]\n" + "\n".join(rows), location=f"table {t_i}"))
    if not res.title:
        for p in d.paragraphs:
            if p.text.strip():
                res.title = p.text.strip()[:200]
                break
    return res


# -------------------------------------------------------------------------- PPTX

def extract_pptx(path: Path) -> ExtractResult:
    from pptx import Presentation

    try:
        prs = Presentation(str(path))
    except Exception as exc:
        raise ExtractionError(f"PPTX could not be opened ({type(exc).__name__})") from exc
    res = ExtractResult(units=[], file_type="pptx")
    res.title = (prs.core_properties.title or "").strip() or None
    res.authors = (prs.core_properties.author or "").strip() or None
    for i, slide in enumerate(prs.slides, start=1):
        parts, heading = [], None
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = "\n".join(p.text for p in shape.text_frame.paragraphs if p.text.strip())
                if t:
                    if heading is None and getattr(shape, "is_placeholder", False) and \
                            "title" in str(shape.placeholder_format.type).lower():
                        heading = t.strip()
                    parts.append(t)
            if getattr(shape, "has_table", False) and shape.has_table:
                rows = [" | ".join(c.text.strip() for c in r.cells) for r in shape.table.rows]
                parts.append("[Table]\n" + "\n".join(rows))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text.strip()
            if notes:
                parts.append("[Speaker notes]\n" + notes)
        text = clean_text("\n\n".join(parts))
        u = Unit(page=i, text=text, location=f"slide {i}", status="ok" if text else "empty")
        if heading:
            u.headings = [(0, heading)]
        if not text:
            u.warnings.append("slide has no extractable text (may contain only images)")
        res.units.append(u)
    if not res.title and res.units and res.units[0].headings:
        res.title = res.units[0].headings[0][1]
    return res


# -------------------------------------------------------------------------- XLSX / CSV

def _rows_to_units(rows: list[list[str]], sheet: str, page: int | None, block: int = 40) -> list[Unit]:
    rows = [r for r in rows if any(c.strip() for c in r)]
    if not rows:
        return [Unit(page=page, text="", status="empty", location=sheet, warnings=["sheet is empty"])]
    header = rows[0]
    units = []
    for start in range(1, max(len(rows), 2), block):
        chunk = rows[start:start + block]
        lines = []
        for r in chunk:
            if len(header) == len(r) and any(h.strip() for h in header):
                lines.append("; ".join(f"{h.strip()}: {v.strip()}" for h, v in zip(header, r) if v.strip()))
            else:
                lines.append(" | ".join(v.strip() for v in r))
        if not chunk:  # header-only sheet
            lines = [" | ".join(header)]
        text = clean_text("\n".join(lines))
        end = start + len(chunk)
        units.append(Unit(page=page, text=text, location=f"{sheet} rows {start + 1}-{max(end, start + 1)}",
                          status="ok" if text else "empty"))
    return units


def extract_xlsx(path: Path) -> ExtractResult:
    import openpyxl

    try:
        wb = openpyxl.load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:
        raise ExtractionError(f"XLSX could not be opened ({type(exc).__name__})") from exc
    res = ExtractResult(units=[], file_type="xlsx")
    for s_i, ws in enumerate(wb.worksheets, start=1):
        rows = [["" if v is None else str(v) for v in row] for row in ws.iter_rows(values_only=True)]
        res.units.extend(_rows_to_units(rows, f"sheet '{ws.title}'", s_i))
    wb.close()
    return res


def _read_text_file(path: Path) -> str:
    raw = path.read_bytes()
    if b"\x00" in raw[:4096]:
        raise ExtractionError("file looks binary, not text")
    for enc in ("utf-8-sig", "cp1256", "latin-1"):
        try:
            return raw.decode(enc)
        except UnicodeDecodeError:
            continue
    raise ExtractionError("unknown text encoding")


def extract_csv(path: Path) -> ExtractResult:
    text = _read_text_file(path)
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    rows = list(csv.reader(io.StringIO(text), dialect))
    res = ExtractResult(units=_rows_to_units(rows, "csv", None), file_type="csv")
    return res


def extract_text(path: Path) -> ExtractResult:
    text = _read_text_file(path)
    is_md = path.suffix.lower() in (".md", ".markdown")
    res = ExtractResult(units=[], file_type="md" if is_md else "txt")
    # Split on markdown headings (or section-name lines for txt).
    lines = text.splitlines()
    sections: list[tuple[str | None, list[str], int]] = []
    heading, buf, start = None, [], 1
    for n, line in enumerate(lines, start=1):
        h = None
        m = re.match(r"^#{1,6}\s+(.+)$", line) if is_md else None
        if m:
            h = m.group(1).strip()
        elif _looks_like_heading(line):
            h = line.strip().rstrip(":")
        if h is not None:
            if any(x.strip() for x in buf):
                sections.append((heading, buf, start))
            heading, buf, start = h, [], n
            if is_md and not res.title and line.startswith("# "):
                res.title = h
            continue
        buf.append(line)
    if any(x.strip() for x in buf):
        sections.append((heading, buf, start))
    for heading, buf, start in sections:
        t = clean_text("\n".join(buf))
        if not t:
            continue
        u = Unit(page=None, text=t, location=f"line {start}")
        if heading:
            u.headings = [(0, heading)]
        res.units.append(u)
    return res
